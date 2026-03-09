"""
Contador de Personas - Servidor Web Multi-Camara
=================================================
Corre headless en Jetson Orin Nano (o cualquier Linux/Windows).
Gestiona 3 camaras USB simultaneas y expone un dashboard web
accesible desde el celular via WiFi local.

Uso:
    python3 server.py

Luego abrir desde el celular:
    http://<ip-jetson>:5000
"""

import cv2
import numpy as np
import supervision as sv
from ultralytics import YOLO
import threading
import time
import csv
import io
import os
import platform
import logging
import signal
from logging.handlers import RotatingFileHandler
from datetime import datetime
from collections import deque
from flask import Flask, jsonify, request, abort, Response

# ---------------------------------------------------------------------------
# Configuracion
# ---------------------------------------------------------------------------
CAMERAS = [
    {"id": 0, "name": "Entrada Principal",   "index": 0},
    {"id": 1, "name": "Entrada Lateral Izq", "index": 1},
    {"id": 2, "name": "Entrada Lateral Der", "index": 2},
]

MODEL_NAME        = "yolov8n.pt"   # mas rapido en Jetson
CONFIDENCE        = 0.35
LINE_POS          = 0.5            # mitad del frame (camara cenital)
LINE_ORIENT       = "horizontal"
COUNTING_MODE     = "line"         # "line" o "fov"
SERVER_PORT       = 5000
AUTOSAVE_INTERVAL = 5 * 60        # segundos

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------
LOG_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "servidor.log")
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        RotatingFileHandler(LOG_PATH, maxBytes=5 * 1024 * 1024, backupCount=3, encoding="utf-8"),
        logging.StreamHandler(),
    ],
)
logger = logging.getLogger("ServidorContador")
logging.getLogger("werkzeug").setLevel(logging.WARNING)

PERSON_CLASS_ID = 0


# ===========================================================================
# PersonCounter
# ===========================================================================
class PersonCounter:
    """Motor de deteccion, tracking y conteo de personas (sin GUI)."""

    def __init__(
        self,
        model_name: str = "yolov8n.pt",
        confidence: float = 0.35,
        line_position: float = 0.5,
        line_orientation: str = "horizontal",
    ):
        self.model_name      = model_name
        self.confidence      = confidence
        self.line_position   = line_position
        self.line_orientation = line_orientation
        self.counting_mode   = "line"

        self.in_count        = 0
        self.out_count       = 0
        self._in_offset      = 0
        self._out_offset     = 0
        self.persons_in_frame = 0
        self.fps             = 0.0

        self.fov_count       = 0
        self._fov_offset     = 0
        self._seen_ids       = set()

        self.interval_data         = []
        self._last_interval_time   = None
        self._interval_in_start    = 0
        self._interval_out_start   = 0
        self._interval_fov_start   = 0
        self.hourly_entries        = {}

        self.model      = None
        self.tracker    = None
        self.line_zone  = None
        self.frame_width  = 0
        self.frame_height = 0
        self._fps_buffer  = deque(maxlen=30)
        self._tracker_reset_time     = time.time()
        self._tracker_reset_interval = 30 * 60

        self.box_annotator   = sv.BoxAnnotator(thickness=2)
        self.label_annotator = sv.LabelAnnotator(text_thickness=1, text_scale=0.5)
        self.trace_annotator = sv.TraceAnnotator(thickness=2, trace_length=60)
        self.line_annotator  = sv.LineZoneAnnotator(thickness=2, text_thickness=2, text_scale=1)

    def load_model(self):
        logger.info("Cargando modelo YOLO: %s", self.model_name)
        self.model = YOLO(self.model_name)
        logger.info("Modelo cargado.")

    def setup_line(self, frame_width: int, frame_height: int):
        self.frame_width  = frame_width
        self.frame_height = frame_height
        self._create_line_zone()
        self._create_tracker()
        if self._last_interval_time is None:
            self._last_interval_time = time.time()

    def _create_line_zone(self):
        if self.line_orientation == "horizontal":
            y = int(self.frame_height * self.line_position)
            start, end = sv.Point(0, y), sv.Point(self.frame_width, y)
        else:
            x = int(self.frame_width * self.line_position)
            start, end = sv.Point(x, 0), sv.Point(x, self.frame_height)
        if self.line_zone is not None:
            self._in_offset  += self.line_zone.in_count
            self._out_offset += self.line_zone.out_count
        self.line_zone = sv.LineZone(
            start=start, end=end,
            triggering_anchors=[sv.Position.TOP_CENTER, sv.Position.BOTTOM_CENTER],
        )

    def _create_tracker(self):
        self.tracker = sv.ByteTrack(
            track_activation_threshold=self.confidence,
            minimum_matching_threshold=0.7,
            lost_track_buffer=45,
            frame_rate=30,
        )
        self._tracker_reset_time = time.time()

    def _periodic_tracker_reset(self):
        logger.info("Reset periodico del tracker.")
        if self.counting_mode == "fov":
            self._fov_offset += len(self._seen_ids)
            self._seen_ids = set()
        self._create_tracker()

    def set_mode(self, mode: str):
        self.counting_mode = mode

    def update_line(self, position: float, orientation: str):
        self.line_position    = position
        self.line_orientation = orientation
        if self.frame_width > 0:
            self._create_line_zone()

    def update_confidence(self, confidence: float):
        self.confidence = confidence

    def process_frame(self, frame: np.ndarray) -> np.ndarray:
        t_start = time.perf_counter()
        if self.model is None:
            return frame

        if time.time() - self._tracker_reset_time > self._tracker_reset_interval:
            self._periodic_tracker_reset()

        results    = self.model(frame, classes=[PERSON_CLASS_ID], conf=self.confidence, verbose=False)[0]
        detections = sv.Detections.from_ultralytics(results)
        detections = self.tracker.update_with_detections(detections)

        if self.counting_mode == "fov":
            self._update_fov_count(detections)
        else:
            self.line_zone.trigger(detections=detections)
            self.in_count  = self._in_offset  + self.line_zone.in_count
            self.out_count = self._out_offset + self.line_zone.out_count

        self.persons_in_frame = len(detections)
        self._check_interval()

        labels = []
        if detections.tracker_id is not None:
            labels = [f"#{tid} {conf:.0%}" for tid, conf in zip(detections.tracker_id, detections.confidence)]

        frame = self.trace_annotator.annotate(scene=frame, detections=detections)
        frame = self.box_annotator.annotate(scene=frame, detections=detections)
        if labels:
            frame = self.label_annotator.annotate(scene=frame, detections=detections, labels=labels)
        if self.counting_mode == "line" and self.line_zone is not None:
            frame = self.line_annotator.annotate(frame=frame, line_counter=self.line_zone)

        if len(detections) > 0 and detections.xyxy is not None:
            for box in detections.xyxy:
                cx, cy = int((box[0]+box[2])/2), int((box[1]+box[3])/2)
                cv2.circle(frame, (cx, cy), 5, (0, 255, 255), -1)
                cv2.circle(frame, (cx, cy), 7, (0, 180, 180), 1)

        elapsed = time.perf_counter() - t_start
        self._fps_buffer.append(1.0 / elapsed if elapsed > 0 else 0)
        self.fps = sum(self._fps_buffer) / len(self._fps_buffer)
        return frame

    def _update_fov_count(self, detections):
        if detections.tracker_id is not None:
            for tid in detections.tracker_id:
                self._seen_ids.add(tid)
        self.fov_count = self._fov_offset + len(self._seen_ids)

    def _check_interval(self):
        now = time.time()
        if self._last_interval_time is None:
            self._last_interval_time = now
            return
        if now - self._last_interval_time >= 15 * 60:
            hora_str = datetime.now().strftime("%H:%M")
            hora     = datetime.now().hour
            if self.counting_mode == "fov":
                fov_interval = self.fov_count - self._interval_fov_start
                self.interval_data.append((hora_str, fov_interval, 0, self.fov_count, 0))
                self.hourly_entries[hora] = self.hourly_entries.get(hora, 0) + fov_interval
                self._interval_fov_start  = self.fov_count
            else:
                in_interval  = self.in_count  - self._interval_in_start
                out_interval = self.out_count - self._interval_out_start
                self.interval_data.append((hora_str, in_interval, out_interval, self.in_count, self.out_count))
                self.hourly_entries[hora]   = self.hourly_entries.get(hora, 0) + in_interval
                self._interval_in_start  = self.in_count
                self._interval_out_start = self.out_count
            self._last_interval_time = now

    def reset_counters(self):
        self.in_count = self.out_count = 0
        self._in_offset = self._out_offset = 0
        self.fov_count = self._fov_offset = 0
        self._seen_ids = set()
        self.persons_in_frame = 0
        self._interval_in_start = self._interval_out_start = self._interval_fov_start = 0
        self._last_interval_time = time.time()
        self.interval_data.clear()
        self.hourly_entries.clear()
        self.line_zone = None
        self._create_line_zone()
        self._create_tracker()
        logger.info("Contadores reiniciados.")

    def export_csv(self, filepath: str):
        data = list(self.interval_data)
        with open(filepath, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            if self.counting_mode == "fov":
                partial = self.fov_count - self._interval_fov_start
                if partial > 0:
                    data.append((datetime.now().strftime("%H:%M"), partial, 0, self.fov_count, 0))
                writer.writerow(["hora", "personas_intervalo", "personas_total"])
                for row in data:
                    writer.writerow([row[0], row[1], row[3]])
            else:
                ip = self.in_count - self._interval_in_start
                op = self.out_count - self._interval_out_start
                if ip > 0 or op > 0:
                    data.append((datetime.now().strftime("%H:%M"), ip, op, self.in_count, self.out_count))
                writer.writerow(["hora", "entradas_intervalo", "salidas_intervalo", "entradas_total", "salidas_total"])
                writer.writerows(data)
        logger.info("CSV exportado: %s", filepath)


# ===========================================================================
# CameraWorker
# ===========================================================================
class CameraWorker:
    """Gestiona una camara USB en un hilo dedicado."""

    WARMUP_ATTEMPTS = 20
    WARMUP_DELAY    = 0.05
    RECONNECT_DELAY = 3

    def __init__(self, cam_id, cam_name, cam_index, model_name,
                 confidence, line_position, line_orientation, counting_mode):
        self.cam_id    = cam_id
        self.cam_name  = cam_name
        self.cam_index = cam_index
        self.connected = False

        self.counter = PersonCounter(
            model_name=model_name,
            confidence=confidence,
            line_position=line_position,
            line_orientation=line_orientation,
        )
        self.counter.set_mode(counting_mode)

        self._lock               = threading.Lock()
        self._stop_event         = threading.Event()
        self._last_annotated_frame = None   # numpy BGR, ultimo frame procesado
        self._pending_cam_index  = None     # si cambia, el loop reabre la camara
        self._stats_cache        = None     # dict publicado tras cada frame; evita bloquear Flask

        self._thread = threading.Thread(
            target=self._loop, daemon=True, name=f"cam-{cam_id}"
        )

    def start(self):
        self._thread.start()
        logger.info("CameraWorker %d (%s) iniciado.", self.cam_id, self.cam_name)

    def stop(self):
        self._stop_event.set()
        self._thread.join(timeout=5)

    # -----------------------------------------------------------------------
    def _build_stats_dict(self) -> dict:
        """Construye el dict de stats. Debe llamarse con self._lock adquirido."""
        c = self.counter
        if c.counting_mode == "fov":
            in_val, out_val, net_val = c.fov_count, 0, c.persons_in_frame
        else:
            in_val  = c.in_count
            out_val = c.out_count
            net_val = c.persons_in_frame
        return {
            "id":        self.cam_id,
            "name":      self.cam_name,
            "in":        in_val,
            "out":       out_val,
            "net":       net_val,
            "in_frame":  c.persons_in_frame,
            "fps":       round(c.fps, 1),
            "connected": self.connected,
            "mode":      c.counting_mode,
            "config": {
                "cam_index":   self.cam_index,
                "line_pos":    round(c.line_position, 2),
                "line_orient": c.line_orientation,
                "mode":        c.counting_mode,
                "confidence":  round(c.confidence, 2),
            },
        }

    def get_stats(self) -> dict:
        """Retorna stats sin bloquear: si YOLO esta corriendo devuelve el cache."""
        if self._lock.acquire(timeout=0.05):
            try:
                stats = self._build_stats_dict()
                self._stats_cache = stats
                return stats
            finally:
                self._lock.release()
        # Lock ocupado (inferencia en curso) — devolver ultimo valor conocido
        if self._stats_cache is not None:
            return self._stats_cache
        # Primera llamada antes del primer frame: esperar
        with self._lock:
            stats = self._build_stats_dict()
            self._stats_cache = stats
            return stats

    def get_snapshot(self) -> bytes | None:
        """Devuelve el ultimo frame anotado como JPEG, o None si no hay frame."""
        with self._lock:
            frame = self._last_annotated_frame
        if frame is None:
            return None
        # Reducir a max 800px ancho para ahorrar ancho de banda
        h, w = frame.shape[:2]
        if w > 800:
            scale = 800 / w
            frame = cv2.resize(frame, (800, int(h * scale)))
        ok, buf = cv2.imencode(".jpg", frame, [cv2.IMWRITE_JPEG_QUALITY, 80])
        return bytes(buf) if ok else None

    def reset(self):
        with self._lock:
            self.counter.reset_counters()
        logger.info("Camara %d reseteada.", self.cam_id)

    def reconfigure(self, cam_index=None, line_pos=None, line_orient=None,
                    mode=None, confidence=None):
        """Aplica nueva configuracion en vivo (thread-safe)."""
        with self._lock:
            c = self.counter
            if mode is not None:
                c.set_mode(mode)
            if confidence is not None:
                c.update_confidence(confidence)
            if line_pos is not None or line_orient is not None:
                pos    = line_pos    if line_pos    is not None else c.line_position
                orient = line_orient if line_orient is not None else c.line_orientation
                c.update_line(pos, orient)
        # Cambio de camara fisica: se maneja fuera del lock principal
        if cam_index is not None and cam_index != self.cam_index:
            self._pending_cam_index = cam_index
            logger.info("Camara %d: pendiente cambio de indice a %d.", self.cam_id, cam_index)

    # -----------------------------------------------------------------------
    def _loop(self):
        cap = None
        while not self._stop_event.is_set():

            # Cambio de camara fisica pendiente
            if self._pending_cam_index is not None:
                new_idx = self._pending_cam_index
                self._pending_cam_index = None
                if cap is not None:
                    cap.release()
                    cap = None
                self.cam_index = new_idx
                with self._lock:
                    self.connected = False
                logger.info("Camara %d: cambiando a indice fisico %d.", self.cam_id, new_idx)

            # Abrir camara si hace falta
            if cap is None or not cap.isOpened():
                cap = self._open_camera(self.cam_index)
                if cap is None:
                    with self._lock:
                        self.connected = False
                    time.sleep(self.RECONNECT_DELAY)
                    continue
                w = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
                h = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
                with self._lock:
                    self.counter.setup_line(w, h)
                    self.connected = True
                logger.info("Camara %d lista (%dx%d).", self.cam_id, w, h)

            ret, frame = cap.read()
            if not ret:
                logger.warning("Camara %d: fallo de lectura. Reconectando...", self.cam_id)
                cap.release()
                cap = None
                with self._lock:
                    self.connected = False
                continue

            try:
                with self._lock:
                    annotated = self.counter.process_frame(frame)
                    self._last_annotated_frame = annotated
            except Exception as e:
                logger.error("Camara %d: error en process_frame: %s", self.cam_id, e)

        if cap is not None:
            cap.release()
        logger.info("Hilo camara %d finalizado.", self.cam_id)

    def _open_camera(self, index: int):
        if platform.system() == "Linux":
            backends = [(cv2.CAP_V4L2, "V4L2"), (cv2.CAP_ANY, "AUTO")]
        else:
            backends = [(cv2.CAP_DSHOW, "DSHOW"), (cv2.CAP_ANY, "AUTO")]

        for backend, name in backends:
            try:
                cap = cv2.VideoCapture(index, backend)
                if not cap.isOpened():
                    cap.release()
                    continue
                cap.set(cv2.CAP_PROP_BUFFERSIZE, 1)
                warmed = False
                for _ in range(self.WARMUP_ATTEMPTS):
                    ret, _ = cap.read()
                    if ret:
                        warmed = True
                        break
                    time.sleep(self.WARMUP_DELAY)
                if not warmed:
                    logger.warning("Camara %d (%s): sin frames en calentamiento.", index, name)
                    cap.release()
                    continue
                logger.info("Camara %d abierta con backend %s.", index, name)
                return cap
            except Exception as e:
                logger.error("Error abriendo camara %d (%s): %s", index, name, e)

        logger.error("No se pudo abrir la camara %d.", index)
        return None


# ===========================================================================
# Dashboard HTML
# ===========================================================================
DASHBOARD_HTML = r"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1,maximum-scale=1">
<title>Contador Personas</title>
<style>
*{box-sizing:border-box;margin:0;padding:0}
body{background:#141414;color:#e0e0e0;font-family:'Segoe UI',system-ui,sans-serif;min-height:100vh;padding:10px}
.brand{text-align:center;margin-bottom:12px;padding:10px 0 8px}
.brand-name{font-size:1.7rem;font-weight:800;letter-spacing:.12em;text-transform:uppercase;background:linear-gradient(90deg,#00e676,#40c4ff);-webkit-background-clip:text;-webkit-text-fill-color:transparent;background-clip:text;line-height:1}
.brand-sub{font-size:.65rem;font-weight:500;color:#555;letter-spacing:.18em;text-transform:uppercase;margin-top:4px}

/* ── Totales ── */
.totals{background:#1e1e1e;border:1px solid #2a2a2a;border-radius:12px;padding:12px 8px;display:flex;justify-content:space-around;align-items:center;margin-bottom:10px}
.t-item{text-align:center;flex:1}
.t-label{font-size:.6rem;color:#555;text-transform:uppercase;letter-spacing:.07em;margin-bottom:3px}
.t-value{font-size:2rem;font-weight:700;line-height:1}
.t-value.in{color:#00e676}.t-value.out{color:#ff5252}.t-value.net{color:#40c4ff}
.t-divider{width:1px;height:36px;background:#2a2a2a;flex-shrink:0}

/* ── Acciones globales ── */
.global-actions{display:flex;gap:8px;margin-bottom:10px}
.btn-danger{flex:1;padding:11px;border-radius:10px;border:none;background:#7f1010;color:#fff;font-size:.82rem;font-weight:600;cursor:pointer}
.btn-danger:active{opacity:.75}

/* ── Grid de camaras ── */
.cameras-grid{display:grid;grid-template-columns:1fr;gap:10px;margin-bottom:10px}
@media(min-width:600px){.cameras-grid{grid-template-columns:repeat(3,1fr)}}

/* ── Card ── */
.cam-card{background:#1e1e1e;border:1px solid #2a2a2a;border-radius:12px;padding:12px}
.cam-card.disconnected{border-color:#ff525230;opacity:.65}
.cam-header{display:flex;justify-content:space-between;align-items:flex-start;margin-bottom:10px}
.cam-name{font-size:.8rem;font-weight:600;color:#bbb;flex:1;padding-right:6px;line-height:1.3}
.badge{font-size:.58rem;font-weight:700;text-transform:uppercase;letter-spacing:.06em;padding:3px 7px;border-radius:20px;flex-shrink:0}
.badge.on{background:#00e67620;color:#00e676;border:1px solid #00e67650}
.badge.off{background:#ff525220;color:#ff5252;border:1px solid #ff525250}

.cam-stats{display:grid;grid-template-columns:1fr 1fr;gap:7px;margin-bottom:8px}
.stat-box{background:#171717;border-radius:8px;padding:9px 8px 7px;text-align:center}
.stat-box.wide{grid-column:1/-1}
.stat-label{font-size:.58rem;color:#555;text-transform:uppercase;letter-spacing:.07em;margin-bottom:3px}
.stat-val{font-size:1.75rem;font-weight:700;line-height:1}
.stat-val.in{color:#00e676}.stat-val.out{color:#ff5252}.stat-val.net{color:#40c4ff}.stat-val.fov{color:#ffab40}

.cam-meta{display:flex;justify-content:space-between;flex-wrap:wrap;gap:4px;font-size:.65rem;color:#444;margin-bottom:9px}
.cam-meta span{white-space:nowrap}

.cam-btns{display:grid;grid-template-columns:1fr 1fr 1fr;gap:6px}
.cam-btns button{padding:8px 0;border-radius:8px;border:1px solid #2e2e2e;background:#222;color:#888;font-size:.72rem;cursor:pointer;transition:background .12s,color .12s}
.cam-btns button:hover{background:#2a2a2a;color:#ccc}
.cam-btns button:active{background:#333}

/* ── Status bar ── */
.status-bar{text-align:center;font-size:.62rem;color:#383838;padding-top:4px}
.status-bar span{color:#484848}

/* ── Modales ── */
.modal{position:fixed;inset:0;z-index:100;display:flex;align-items:flex-end;justify-content:center}
.modal.hidden{display:none}
.modal-overlay{position:absolute;inset:0;background:#000a}
.modal-box{position:relative;width:100%;max-width:520px;max-height:92vh;background:#1a1a1a;border-radius:18px 18px 0 0;display:flex;flex-direction:column;overflow:hidden}
@media(min-width:600px){.modal{align-items:center}.modal-box{border-radius:16px;margin:16px;max-height:88vh}}

.modal-header{display:flex;justify-content:space-between;align-items:center;padding:14px 16px 12px;border-bottom:1px solid #252525;flex-shrink:0}
.modal-header .title{font-size:.88rem;font-weight:600;color:#ccc}
.modal-close{background:none;border:none;color:#555;font-size:1.1rem;cursor:pointer;padding:4px 8px;border-radius:6px}
.modal-close:hover{color:#aaa;background:#252525}
.modal-body{overflow-y:auto;padding:14px 16px;flex:1}
.modal-footer{display:flex;gap:8px;padding:12px 16px;border-top:1px solid #252525;flex-shrink:0}
.modal-footer button{flex:1;padding:11px;border-radius:9px;border:none;font-size:.82rem;font-weight:600;cursor:pointer}
.btn-secondary{background:#252525;color:#888}
.btn-primary{background:#1a5c2e;color:#00e676}
.btn-primary:active,.btn-secondary:active{opacity:.75}

/* ── Snapshot modal content ── */
#snap-img{width:100%;border-radius:8px;display:block}
#snap-loading,#snap-error{text-align:center;padding:32px;font-size:.82rem;color:#555}
#snap-error{color:#ff5252}
#snap-time{font-size:.68rem;color:#444;text-align:center;margin-top:6px}

/* ── Config form ── */
.form-group{margin-bottom:14px}
.form-group label.field-label{display:block;font-size:.72rem;color:#888;margin-bottom:4px;text-transform:uppercase;letter-spacing:.06em}
.form-hint{font-size:.65rem;color:#444;margin-bottom:5px}
input[type=number]{width:100%;background:#111;border:1px solid #2e2e2e;border-radius:7px;padding:9px 10px;color:#ddd;font-size:.9rem}
input[type=range]{width:100%;accent-color:#00e676;margin-top:4px}
.range-val{font-size:.78rem;color:#aaa;margin-top:2px;text-align:right}

.radio-cards{display:grid;grid-template-columns:1fr 1fr;gap:7px;margin-top:4px}
.radio-card{position:relative}
.radio-card input{position:absolute;opacity:0;width:0;height:0}
.radio-card span{display:block;border:1px solid #2e2e2e;border-radius:9px;padding:10px 8px;text-align:center;cursor:pointer;transition:border-color .15s,background .15s}
.radio-card span .rc-title{font-size:.78rem;font-weight:600;color:#aaa;display:block;margin-bottom:2px}
.radio-card span .rc-desc{font-size:.62rem;color:#555;line-height:1.3}
.radio-card input:checked + span{border-color:#00e676;background:#00e67610}
.radio-card input:checked + span .rc-title{color:#00e676}

.toggle-group{display:grid;grid-template-columns:1fr 1fr;gap:7px;margin-top:4px}
.toggle-opt{position:relative}
.toggle-opt input{position:absolute;opacity:0;width:0;height:0}
.toggle-opt span{display:block;border:1px solid #2e2e2e;border-radius:8px;padding:8px;text-align:center;font-size:.78rem;color:#666;cursor:pointer;transition:border-color .15s,color .15s,background .15s}
.toggle-opt input:checked + span{border-color:#40c4ff;color:#40c4ff;background:#40c4ff10}

.line-controls-wrapper{transition:opacity .2s}
.line-controls-wrapper.dimmed{opacity:.3;pointer-events:none}

/* ── Botones extra ── */
.btn-chart{flex:1;padding:11px;border-radius:10px;border:none;background:#0d2a45;color:#40c4ff;font-size:.82rem;font-weight:600;cursor:pointer}
.btn-export{flex:1;padding:11px;border-radius:10px;border:none;background:#0d2d1a;color:#00e676;font-size:.82rem;font-weight:600;cursor:pointer}
.btn-chart:active,.btn-export:active{opacity:.75}

/* ── Grafica horaria ── */
#hourly-canvas{width:100%;height:auto;border-radius:8px;display:block;background:#161616}
.chart-legend{display:flex;flex-wrap:wrap;gap:10px;margin-top:12px;justify-content:center}
.legend-item{display:flex;align-items:center;gap:6px;font-size:.72rem;color:#888}
.legend-swatch{width:12px;height:12px;border-radius:3px;flex-shrink:0}
</style>
</head>
<body>
<div class="brand">
  <div class="brand-name">Xenith AI</div>
  <div class="brand-sub">Sistema de conteo de personas</div>
</div>

<div class="totals">
  <div class="t-item">
    <div class="t-label">Total Entradas</div>
    <div class="t-value in" id="total-in">--</div>
  </div>
  <div class="t-divider"></div>
  <div class="t-item">
    <div class="t-label">Total Salidas</div>
    <div class="t-value out" id="total-out">--</div>
  </div>
  <div class="t-divider"></div>
  <div class="t-item">
    <div class="t-label">En Stand</div>
    <div class="t-value net" id="total-net">--</div>
  </div>
</div>

<div class="global-actions">
  <button class="btn-danger" onclick="resetAll()">&#8635; Reiniciar</button>
  <button class="btn-chart"  onclick="openChart()">&#128202; Gr&#225;fica</button>
  <button class="btn-export" onclick="downloadExcel()">&#8659; Excel</button>
</div>

<div class="cameras-grid" id="cameras-grid">
  <div style="color:#333;text-align:center;padding:40px;grid-column:1/-1">Cargando...</div>
</div>

<div class="status-bar">Actualizado: <span id="last-update">--</span></div>

<!-- ══════════════════════════════════════════════════
     MODAL: SNAPSHOT
═══════════════════════════════════════════════════ -->
<div id="snapshot-modal" class="modal hidden">
  <div class="modal-overlay" onclick="closeModal('snapshot')"></div>
  <div class="modal-box">
    <div class="modal-header">
      <span class="title" id="snap-title">Frame</span>
      <button class="modal-close" onclick="closeModal('snapshot')">&#10005;</button>
    </div>
    <div class="modal-body">
      <div id="snap-loading">Cargando frame...</div>
      <div id="snap-error" style="display:none"></div>
      <img id="snap-img" style="display:none" alt="Snapshot">
      <div id="snap-time"></div>
    </div>
    <div class="modal-footer">
      <button class="btn-secondary" onclick="closeModal('snapshot')">Cerrar</button>
      <button class="btn-primary" onclick="refreshSnapshot()">&#8635; Actualizar frame</button>
    </div>
  </div>
</div>

<!-- ══════════════════════════════════════════════════
     MODAL: CONFIG
═══════════════════════════════════════════════════ -->
<div id="config-modal" class="modal hidden">
  <div class="modal-overlay" onclick="closeModal('config')"></div>
  <div class="modal-box">
    <div class="modal-header">
      <span class="title" id="cfg-title">Configurar camara</span>
      <button class="modal-close" onclick="closeModal('config')">&#10005;</button>
    </div>
    <div class="modal-body">

      <div class="form-group">
        <label class="field-label">Camara fisica (indice)</label>
        <div class="form-hint">Numero del dispositivo: 0, 1, 2... Cambiar reconecta la camara.</div>
        <input type="number" id="cfg-cam-index" min="0" max="10" step="1" value="0">
      </div>

      <div class="form-group">
        <label class="field-label">Modo de conteo</label>
        <div class="radio-cards">
          <label class="radio-card">
            <input type="radio" name="cfg-mode" value="line" onchange="updateModeUI()">
            <span><span class="rc-title">Cruce de linea</span><span class="rc-desc">Cuenta entradas y salidas al cruzar</span></span>
          </label>
          <label class="radio-card">
            <input type="radio" name="cfg-mode" value="fov" onchange="updateModeUI()">
            <span><span class="rc-title">Campo de vision</span><span class="rc-desc">Cuenta personas unicas vistas</span></span>
          </label>
        </div>
      </div>

      <div class="line-controls-wrapper" id="line-controls">
        <div class="form-group">
          <label class="field-label">Posicion de linea: <span id="cfg-line-pos-val">50%</span></label>
          <input type="range" id="cfg-line-pos" min="0" max="1" step="0.05"
                 oninput="document.getElementById('cfg-line-pos-val').textContent=Math.round(this.value*100)+'%'">
        </div>
        <div class="form-group">
          <label class="field-label">Orientacion de linea</label>
          <div class="toggle-group">
            <label class="toggle-opt">
              <input type="radio" name="cfg-orient" value="horizontal">
              <span>Horizontal &mdash;</span>
            </label>
            <label class="toggle-opt">
              <input type="radio" name="cfg-orient" value="vertical">
              <span>Vertical &#124;</span>
            </label>
          </div>
        </div>
      </div>

      <div class="form-group">
        <label class="field-label">Umbral de confianza: <span id="cfg-conf-val">0.35</span></label>
        <input type="range" id="cfg-confidence" min="0.1" max="0.9" step="0.05"
               oninput="document.getElementById('cfg-conf-val').textContent=parseFloat(this.value).toFixed(2)">
      </div>

    </div>
    <div class="modal-footer">
      <button class="btn-secondary" onclick="closeModal('config')">Cancelar</button>
      <button class="btn-primary" id="save-cfg-btn" onclick="saveConfig()">Guardar cambios</button>
    </div>
  </div>
</div>

<!-- ══════════════════════════════════════════════════
     MODAL: GRAFICA HORARIA
═══════════════════════════════════════════════════ -->
<div id="chart-modal" class="modal hidden">
  <div class="modal-overlay" onclick="closeModal('chart')"></div>
  <div class="modal-box">
    <div class="modal-header">
      <span class="title">Personas por hora del d&#237;a</span>
      <button class="modal-close" onclick="closeModal('chart')">&#10005;</button>
    </div>
    <div class="modal-body">
      <canvas id="hourly-canvas" width="700" height="240"></canvas>
      <div class="chart-legend" id="chart-legend"></div>
    </div>
    <div class="modal-footer">
      <button class="btn-secondary" onclick="closeModal('chart')">Cerrar</button>
      <button class="btn-primary" onclick="loadChart()">&#8635; Actualizar</button>
    </div>
  </div>
</div>

<script>
'use strict';

let activeCamId  = null;
let lastCamData  = {};   // id -> cam stats+config

// ─── helpers ──────────────────────────────────────────────────────
function setVal(id, val) {
  const el = document.getElementById(id);
  if (el && el.textContent !== String(val)) el.textContent = val;
}

function closeModal(type) {
  document.getElementById(type + '-modal').classList.add('hidden');
  if (type === 'snapshot') activeCamId = null;
}

// ─── build camera card (first render) ─────────────────────────────
function buildCard(cam) {
  const line = cam.mode !== 'fov';
  const conn = cam.connected;
  const cfg  = cam.config;
  const modeLabel = cam.mode === 'fov' ? 'FOV' : 'Linea';

  let statsHtml;
  if (line) {
    statsHtml = `
      <div class="stat-box">
        <div class="stat-label">Entradas</div>
        <div class="stat-val in" id="c${cam.id}-in">${cam.in}</div>
      </div>
      <div class="stat-box">
        <div class="stat-label">Salidas</div>
        <div class="stat-val out" id="c${cam.id}-out">${cam.out}</div>
      </div>
      <div class="stat-box wide">
        <div class="stat-label">En Stand</div>
        <div class="stat-val net" id="c${cam.id}-net">${cam.net}</div>
      </div>`;
  } else {
    statsHtml = `
      <div class="stat-box wide">
        <div class="stat-label">Personas vistas</div>
        <div class="stat-val fov" id="c${cam.id}-in">${cam.in}</div>
      </div>`;
  }

  return `<div class="cam-card${conn ? '' : ' disconnected'}" id="card-${cam.id}">
    <div class="cam-header">
      <div class="cam-name">${cam.name}</div>
      <span class="badge ${conn ? 'on' : 'off'}" id="c${cam.id}-badge">${conn ? 'Online' : 'Offline'}</span>
    </div>
    <div class="cam-stats">${statsHtml}</div>
    <div class="cam-meta">
      <span>En cuadro: <b id="c${cam.id}-frame">${cam.in_frame}</b></span>
      <span>FPS: <b id="c${cam.id}-fps">${cam.fps}</b></span>
      <span>Modo: <b id="c${cam.id}-mode">${modeLabel}</b></span>
      <span>Cam #<b id="c${cam.id}-idx">${cfg.cam_index}</b></span>
    </div>
    <div class="cam-btns">
      <button onclick="openSnapshot(${cam.id})">&#128247; Frame</button>
      <button onclick="openConfig(${cam.id})">&#9881;&#65039; Config</button>
      <button onclick="resetCam(${cam.id})">&#8635; Reset</button>
    </div>
  </div>`;
}

// ─── patch existing card (subsequent updates) ──────────────────────
function patchCard(cam) {
  const card = document.getElementById('card-' + cam.id);
  if (!card) return false;

  const conn = cam.connected;
  card.className = 'cam-card' + (conn ? '' : ' disconnected');

  const badge = document.getElementById('c' + cam.id + '-badge');
  if (badge) { badge.textContent = conn ? 'Online' : 'Offline'; badge.className = 'badge ' + (conn ? 'on' : 'off'); }

  setVal('c' + cam.id + '-in',    cam.in);
  setVal('c' + cam.id + '-out',   cam.out);
  setVal('c' + cam.id + '-net',   cam.net);
  setVal('c' + cam.id + '-frame', cam.in_frame);
  setVal('c' + cam.id + '-fps',   cam.fps);
  setVal('c' + cam.id + '-mode',  cam.mode === 'fov' ? 'FOV' : 'Linea');
  setVal('c' + cam.id + '-idx',   cam.config.cam_index);
  return true;
}

// ─── polling ──────────────────────────────────────────────────────
async function fetchCounts() {
  try {
    const resp = await fetch('/api/counts');
    if (!resp.ok) return;
    const data = await resp.json();

    setVal('total-in',  data.totals.in);
    setVal('total-out', data.totals.out);
    setVal('total-net', data.totals.net);

    const grid = document.getElementById('cameras-grid');
    let needBuild = false;

    for (const cam of data.cameras) {
      lastCamData[cam.id] = cam;
      if (!patchCard(cam)) { needBuild = true; break; }
    }

    if (needBuild) {
      grid.innerHTML = data.cameras.map(buildCard).join('');
    }

    document.getElementById('last-update').textContent = data.timestamp;
  } catch(e) { /* network error - silently ignore */ }
}

// ─── snapshot ─────────────────────────────────────────────────────
function openSnapshot(camId) {
  activeCamId = camId;
  const cam = lastCamData[camId];
  document.getElementById('snap-title').textContent = cam ? cam.name : 'Camara ' + camId;
  document.getElementById('snap-img').style.display = 'none';
  document.getElementById('snap-error').style.display = 'none';
  document.getElementById('snap-loading').style.display = 'block';
  document.getElementById('snap-time').textContent = '';
  document.getElementById('snapshot-modal').classList.remove('hidden');
  refreshSnapshot();
}

async function refreshSnapshot() {
  if (activeCamId === null) return;
  const img     = document.getElementById('snap-img');
  const loading = document.getElementById('snap-loading');
  const errEl   = document.getElementById('snap-error');

  img.style.display = 'none';
  errEl.style.display = 'none';
  loading.style.display = 'block';

  try {
    const resp = await fetch('/api/snapshot/' + activeCamId + '?t=' + Date.now());
    if (resp.ok) {
      const blob = await resp.blob();
      img.src = URL.createObjectURL(blob);
      img.style.display = 'block';
      document.getElementById('snap-time').textContent = 'Tomado: ' + new Date().toLocaleTimeString();
    } else if (resp.status === 503) {
      errEl.textContent = 'Camara no disponible o sin frames aun. Intenta en unos segundos.';
      errEl.style.display = 'block';
    } else {
      errEl.textContent = 'Error al obtener el frame (' + resp.status + ')';
      errEl.style.display = 'block';
    }
  } catch(e) {
    errEl.textContent = 'Error de conexion';
    errEl.style.display = 'block';
  } finally {
    loading.style.display = 'none';
  }
}

// ─── config ───────────────────────────────────────────────────────
function openConfig(camId) {
  activeCamId = camId;
  const cam = lastCamData[camId];
  if (!cam) return;
  const cfg = cam.config;

  document.getElementById('cfg-title').textContent = 'Config: ' + cam.name;
  document.getElementById('cfg-cam-index').value = cfg.cam_index;
  document.getElementById('cfg-line-pos').value  = cfg.line_pos;
  document.getElementById('cfg-line-pos-val').textContent = Math.round(cfg.line_pos * 100) + '%';
  document.getElementById('cfg-confidence').value = cfg.confidence;
  document.getElementById('cfg-conf-val').textContent = parseFloat(cfg.confidence).toFixed(2);

  const modeEl = document.querySelector(`input[name="cfg-mode"][value="${cfg.mode}"]`);
  if (modeEl) modeEl.checked = true;

  const orientEl = document.querySelector(`input[name="cfg-orient"][value="${cfg.line_orient}"]`);
  if (orientEl) orientEl.checked = true;

  updateModeUI();
  document.getElementById('config-modal').classList.remove('hidden');
}

function updateModeUI() {
  const modeEl = document.querySelector('input[name="cfg-mode"]:checked');
  const wrapper = document.getElementById('line-controls');
  if (!wrapper) return;
  const isLine = modeEl && modeEl.value === 'line';
  wrapper.classList.toggle('dimmed', !isLine);
}

async function saveConfig() {
  if (activeCamId === null) return;

  const modeEl   = document.querySelector('input[name="cfg-mode"]:checked');
  const orientEl = document.querySelector('input[name="cfg-orient"]:checked');

  const payload = {
    cam_index:   parseInt(document.getElementById('cfg-cam-index').value),
    line_pos:    parseFloat(document.getElementById('cfg-line-pos').value),
    line_orient: orientEl ? orientEl.value : 'horizontal',
    mode:        modeEl  ? modeEl.value  : 'line',
    confidence:  parseFloat(document.getElementById('cfg-confidence').value),
  };

  const btn = document.getElementById('save-cfg-btn');
  btn.textContent = 'Guardando...';
  btn.disabled = true;

  try {
    const resp = await fetch('/api/config/' + activeCamId, {
      method: 'POST',
      headers: {'Content-Type': 'application/json'},
      body: JSON.stringify(payload),
    });
    if (resp.ok) closeModal('config');
    else { alert('Error al guardar config: ' + resp.status); }
  } catch(e) {
    alert('Error de conexion al guardar.');
  } finally {
    btn.textContent = 'Guardar cambios';
    btn.disabled = false;
  }
}

// ─── reset ────────────────────────────────────────────────────────
async function resetCam(id) {
  if (!confirm('Reiniciar contadores de esta camara?')) return;
  await fetch('/api/reset/' + id, {method:'POST'}).catch(()=>{});
}

async function resetAll() {
  if (!confirm('Reiniciar contadores de TODAS las camaras?')) return;
  await fetch('/api/reset/all', {method:'POST'}).catch(()=>{});
}

// ─── chart ────────────────────────────────────────────────────────
const CAM_COLORS = ['#00e676', '#40c4ff', '#ffab40'];

async function openChart() {
  document.getElementById('chart-modal').classList.remove('hidden');
  await loadChart();
}

async function loadChart() {
  try {
    const resp = await fetch('/api/stats');
    if (!resp.ok) return;
    const data = await resp.json();
    renderChart(data);
    renderLegend(data.cameras);
  } catch(e) {}
}

function renderLegend(cameras) {
  document.getElementById('chart-legend').innerHTML = cameras.map((cam, i) =>
    `<div class="legend-item">
       <div class="legend-swatch" style="background:${CAM_COLORS[i]||'#888'}"></div>
       ${cam.name}
     </div>`
  ).join('');
}

function renderChart(data) {
  const canvas  = document.getElementById('hourly-canvas');
  const ctx     = canvas.getContext('2d');
  const cameras = data.cameras;
  const W = 700, H = 240;
  const pad = {left: 40, right: 8, top: 18, bottom: 32};
  const cW = W - pad.left - pad.right;
  const cH = H - pad.top  - pad.bottom;
  const bw = cW / 24;

  // valores por camara por hora
  const vals = cameras.map(cam =>
    Array.from({length: 24}, (_, h) => cam.hourly[String(h)] || 0)
  );
  const hourTotals = Array.from({length: 24}, (_, h) =>
    cameras.reduce((s, _, ci) => s + vals[ci][h], 0)
  );
  const maxV = Math.max(...hourTotals, 1);

  ctx.clearRect(0, 0, W, H);
  ctx.fillStyle = '#161616';
  ctx.fillRect(0, 0, W, H);

  // lineas de cuadricula
  for (let i = 0; i <= 4; i++) {
    const y = pad.top + cH - (i / 4) * cH;
    ctx.strokeStyle = '#252525';
    ctx.lineWidth = 1;
    ctx.beginPath(); ctx.moveTo(pad.left, y); ctx.lineTo(pad.left + cW, y); ctx.stroke();
    ctx.fillStyle = '#555';
    ctx.font = '10px sans-serif';
    ctx.textAlign = 'right';
    ctx.fillText(Math.round(maxV * i / 4), pad.left - 5, y + 3);
  }

  // barras apiladas por camara
  for (let h = 0; h < 24; h++) {
    let base = pad.top + cH;
    const x  = pad.left + h * bw + 1;
    cameras.forEach((_, ci) => {
      const v = vals[ci][h];
      if (!v) return;
      const bh = (v / maxV) * cH;
      base -= bh;
      ctx.fillStyle = CAM_COLORS[ci] || '#888';
      ctx.fillRect(x, base, bw - 2, bh);
    });
  }

  // resalte hora actual
  const nowH = new Date().getHours();
  ctx.fillStyle = '#ffffff0a';
  ctx.fillRect(pad.left + nowH * bw, pad.top, bw, cH);

  // etiquetas eje X (cada 2 horas)
  for (let h = 0; h < 24; h += 2) {
    ctx.fillStyle = '#555';
    ctx.font = '9px sans-serif';
    ctx.textAlign = 'center';
    ctx.fillText(h + 'h', pad.left + h * bw + bw / 2, H - 10);
  }
}

// ─── excel ────────────────────────────────────────────────────────
function downloadExcel() {
  window.location.href = '/api/export';
}

// ─── start ────────────────────────────────────────────────────────
fetchCounts();
setInterval(fetchCounts, 1000);
</script>
</body>
</html>
"""


# ===========================================================================
# Flask App
# ===========================================================================
app = Flask(__name__)

camera_workers: list[CameraWorker] = []


@app.route("/")
def dashboard():
    return DASHBOARD_HTML, 200, {"Content-Type": "text/html; charset=utf-8"}


@app.route("/api/counts")
def api_counts():
    cameras_data = [w.get_stats() for w in camera_workers]
    return jsonify({
        "timestamp": datetime.now().strftime("%H:%M:%S"),
        "cameras":   cameras_data,
        "totals": {
            "in":  sum(c["in"]  for c in cameras_data),
            "out": sum(c["out"] for c in cameras_data),
            "net": sum(c["net"] for c in cameras_data),
        },
    })


@app.route("/api/snapshot/<int:cam_id>")
def api_snapshot(cam_id):
    for w in camera_workers:
        if w.cam_id == cam_id:
            data = w.get_snapshot()
            if data is None:
                abort(503)
            return Response(data, mimetype="image/jpeg")
    abort(404)


@app.route("/api/config/<int:cam_id>", methods=["POST"])
def api_config(cam_id):
    body = request.get_json(force=True, silent=True) or {}
    for w in camera_workers:
        if w.cam_id == cam_id:
            w.reconfigure(
                cam_index   = body.get("cam_index"),
                line_pos    = body.get("line_pos"),
                line_orient = body.get("line_orient"),
                mode        = body.get("mode"),
                confidence  = body.get("confidence"),
            )
            logger.info("Config camara %d actualizada: %s", cam_id, body)
            return jsonify({"status": "ok"})
    abort(404)


@app.route("/api/reset/<cam_id>", methods=["POST"])
def api_reset(cam_id):
    if cam_id == "all":
        for w in camera_workers:
            w.reset()
        return jsonify({"status": "ok", "reset": "all"})
    try:
        cid = int(cam_id)
    except ValueError:
        abort(400)
    for w in camera_workers:
        if w.cam_id == cid:
            w.reset()
            return jsonify({"status": "ok", "reset": cid})
    abort(404)


@app.route("/api/reset/all", methods=["POST"])
def api_reset_all():
    for w in camera_workers:
        w.reset()
    return jsonify({"status": "ok", "reset": "all"})


# ===========================================================================
# Stats horarias y exportacion Excel
# ===========================================================================
@app.route("/api/stats")
def api_stats():
    cameras = []
    for w in camera_workers:
        with w._lock:
            c = w.counter
            hourly = dict(c.hourly_entries)
            # incluir intervalo parcial actual (aun no volcado a hourly_entries)
            current_hour = datetime.now().hour
            if c.counting_mode == "fov":
                partial = c.fov_count - c._interval_fov_start
            else:
                partial = c.in_count - c._interval_in_start
            if partial > 0:
                hourly[current_hour] = hourly.get(current_hour, 0) + partial
        cameras.append({
            "id":     w.cam_id,
            "name":   w.cam_name,
            "hourly": {str(k): v for k, v in hourly.items()},
        })
    return jsonify({"cameras": cameras})


@app.route("/api/export")
def api_export():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return "Dependencia faltante. Ejecuta: pip install openpyxl", 501

    date_str = datetime.now().strftime("%Y-%m-%d")

    # snapshot thread-safe de todos los contadores
    snaps = []
    for w in camera_workers:
        with w._lock:
            c = w.counter
            intervals = list(c.interval_data)
            # agregar intervalo parcial actual
            if c.counting_mode == "fov":
                partial_in = c.fov_count - c._interval_fov_start
                partial_out = 0
                if partial_in > 0:
                    intervals.append((datetime.now().strftime("%H:%M"), partial_in, 0, c.fov_count, 0))
            else:
                partial_in  = c.in_count  - c._interval_in_start
                partial_out = c.out_count - c._interval_out_start
                if partial_in > 0 or partial_out > 0:
                    intervals.append((datetime.now().strftime("%H:%M"), partial_in, partial_out, c.in_count, c.out_count))
            hourly = dict(c.hourly_entries)
            ch = datetime.now().hour
            p = partial_in if c.counting_mode != "fov" else (c.fov_count - c._interval_fov_start)
            if p > 0:
                hourly[ch] = hourly.get(ch, 0) + p
            snaps.append({
                "id":        w.cam_id,
                "name":      w.cam_name,
                "mode":      c.counting_mode,
                "intervals": intervals,
                "hourly":    hourly,
            })

    wb = Workbook()

    # ── Hoja 1: resumen por hora ──────────────────────────────────────
    ws = wb.active
    ws.title = "Resumen por hora"
    header = ["Hora"] + [s["name"] for s in snaps] + ["TOTAL"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 8
    for col_idx in range(2, len(snaps) + 3):
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[col_letter].width = 22

    grand_total = 0
    for hour in range(24):
        row = [f"{hour:02d}:00"]
        row_sum = 0
        for s in snaps:
            val = s["hourly"].get(hour, 0)
            row.append(val)
            row_sum += val
        row.append(row_sum)
        grand_total += row_sum
        ws.append(row)

    total_row = ["TOTAL"] + [sum(s["hourly"].values()) for s in snaps] + [grand_total]
    ws.append(total_row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    # ── Hoja por camara ───────────────────────────────────────────────
    for s in snaps:
        ws2 = wb.create_sheet(title=s["name"][:31])
        if s["mode"] == "fov":
            ws2.append(["Hora", "Personas intervalo", "Personas total"])
            for row in s["intervals"]:
                ws2.append([row[0], row[1], row[3]])
        else:
            ws2.append(["Hora", "Entradas intervalo", "Salidas intervalo",
                        "Entradas total", "Salidas total"])
            for row in s["intervals"]:
                ws2.append([row[0], row[1], row[2], row[3], row[4]])
        for cell in ws2[1]:
            cell.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=conteo_{date_str}.xlsx"},
    )


# ===========================================================================
# Auto-guardado CSV
# ===========================================================================
def _autosave_loop():
    csv_dir = os.path.dirname(os.path.abspath(__file__))
    while True:
        time.sleep(AUTOSAVE_INTERVAL)
        date_str = datetime.now().strftime("%Y-%m-%d")
        for w in camera_workers:
            stats = w.get_stats()
            if stats["in"] > 0 or stats["out"] > 0:
                filepath = os.path.join(csv_dir, f"conteo_cam{w.cam_id}_{date_str}.csv")
                try:
                    with w._lock:
                        w.counter.export_csv(filepath)
                except Exception as e:
                    logger.error("Error auto-guardado camara %d: %s", w.cam_id, e)


# ===========================================================================
# Watchdog de hilos de camara
# ===========================================================================
def _watchdog_loop():
    """Revisa cada 30 s que los hilos de camara sigan vivos; los reinicia si no."""
    time.sleep(60)  # dar tiempo al arranque
    while True:
        time.sleep(30)
        for w in camera_workers:
            if not w._thread.is_alive() and not w._stop_event.is_set():
                logger.error(
                    "Watchdog: hilo camara %d (%s) muerto. Reiniciando...",
                    w.cam_id, w.cam_name,
                )
                w._thread = threading.Thread(
                    target=w._loop, daemon=True, name=f"cam-{w.cam_id}"
                )
                w._thread.start()


# ===========================================================================
# Cierre limpio (guarda CSV al recibir SIGTERM / Ctrl-C)
# ===========================================================================
def _save_all_csv():
    csv_dir  = os.path.dirname(os.path.abspath(__file__))
    date_str = datetime.now().strftime("%Y-%m-%d")
    for w in camera_workers:
        filepath = os.path.join(csv_dir, f"conteo_cam{w.cam_id}_{date_str}.csv")
        try:
            with w._lock:
                w.counter.export_csv(filepath)
            logger.info("CSV guardado: %s", filepath)
        except Exception as e:
            logger.error("Error guardando CSV camara %d: %s", w.cam_id, e)


def _shutdown(signum, frame):
    logger.info("Senal %s recibida. Guardando datos y cerrando...", signum)
    _save_all_csv()
    for w in camera_workers:
        w.stop()
    logger.info("Servidor detenido limpiamente.")
    os._exit(0)


# ===========================================================================
# Main
# ===========================================================================
def main():
    global camera_workers

    logger.info("=" * 60)
    logger.info("Servidor Multi-Camara - Contador de Personas")
    logger.info("Modelo: %s | Confianza: %.2f | Modo: %s", MODEL_NAME, CONFIDENCE, COUNTING_MODE)
    logger.info("=" * 60)

    workers = []
    for cam_cfg in CAMERAS:
        w = CameraWorker(
            cam_id         = cam_cfg["id"],
            cam_name       = cam_cfg["name"],
            cam_index      = cam_cfg["index"],
            model_name     = MODEL_NAME,
            confidence     = CONFIDENCE,
            line_position  = LINE_POS,
            line_orientation = LINE_ORIENT,
            counting_mode  = COUNTING_MODE,
        )
        logger.info("Cargando modelo para camara %d (%s)...", cam_cfg["id"], cam_cfg["name"])
        w.counter.load_model()
        workers.append(w)

    camera_workers = workers

    # Registrar cierre limpio
    signal.signal(signal.SIGTERM, _shutdown)
    signal.signal(signal.SIGINT,  _shutdown)

    for w in camera_workers:
        w.start()

    threading.Thread(target=_autosave_loop, daemon=True, name="autosave").start()
    threading.Thread(target=_watchdog_loop, daemon=True, name="watchdog").start()

    logger.info("Servidor en http://0.0.0.0:%d", SERVER_PORT)
    logger.info("Acceder desde la red: http://<ip-dispositivo>:%d", SERVER_PORT)

    app.run(host="0.0.0.0", port=SERVER_PORT, threaded=True, use_reloader=False)


if __name__ == "__main__":
    main()
